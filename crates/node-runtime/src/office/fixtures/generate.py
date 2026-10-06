"""Generate owned document fixtures with an external test environment."""
from pathlib import Path
import json
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen.canvas import Canvas
import xlsxwriter

root = Path(__file__).parent
root.mkdir(parents=True, exist_ok=True)
doc = Document()
style = doc.styles['Normal']
style.font.name = 'Noto Sans SC'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans SC')
doc.add_heading('文档验收', 0)
doc.add_paragraph('中文内容与 English text')
table = doc.add_table(rows=2, cols=2)
for row, values in zip(table.rows, [('项目', '数量'), ('报告', '12')]):
    for cell, value in zip(row.cells, values):
        cell.text = value
doc.add_page_break()
doc.add_paragraph('Second page')
doc.save(root / 'report.docx')

book = xlsxwriter.Workbook(root / 'book.xlsx')
style = book.add_format({'font_name': 'Noto Sans SC'})
sheet = book.add_worksheet('预算')
sheet.set_column('A:B', 20, style)
sheet.write_row('A1', ['数量', '合计'])
sheet.write_column('A2', [12, 8])
sheet.write_formula('B2', '=SUM(A2:A3)', style, 20)
book.add_worksheet('备注').write('A1', 'Second sheet')
book.close()

slides = Presentation()
for title in ['中文幻灯片', 'Second slide']:
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame
    run = text.paragraphs[0].add_run()
    run.text = title
    run.font.name = 'Noto Sans SC'
    slide.notes_slide.notes_text_frame.text = 'Speaker notes'
slides.save(root / 'deck.pptx')

pdf = Canvas(str(root / 'report.pdf'), invariant=1)
pdf.drawString(72, 750, 'PDF acceptance')
pdf.showPage()
pdf.drawString(72, 750, 'Second page')
pdf.save()

manifest = [
    {'path': 'report.docx', 'text': '中文内容', 'pages': 2},
    {'path': 'report.pdf', 'text': 'PDF acceptance', 'pages': 2},
    {'path': 'book.xlsx', 'text': '20', 'pages': 2},
    {'path': 'deck.pptx', 'text': '中文幻灯片', 'pages': 2},
]
(root.parent / 'fixtures.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')
