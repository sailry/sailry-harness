"""Exercise real authoring APIs and reopen saved outputs without system Python packages."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt as SlidePt
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from pypdf import PdfReader, PdfWriter
import pypdfium2 as pdfium
import pdfplumber
import xlsxwriter

import os
output = Path('output/office-check')
output.mkdir(parents=True, exist_ok=True)
os.chdir(output)

# Exercise typography, page geometry, fields and styles with the complete library API.
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin, section.bottom_margin = Mm(30), Mm(25)
section.left_margin = section.right_margin = Mm(28)
section.header.paragraphs[0].text = 'Office acceptance'
p = doc.add_paragraph()
p.paragraph_format.line_spacing = Pt(28)
p.paragraph_format.space_after = Pt(12)
run = p.add_run('Administrative document')
run.font.name, run.font.size = 'Noto Sans CJK SC', Pt(22)
run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
run.font.color.rgb = RGBColor.from_string('C00000')
borders = OxmlElement('w:pBdr')
line = OxmlElement('w:bottom')
for name, value in {'val':'single', 'sz':'12', 'space':'4', 'color':'C00000'}.items():
    line.set(qn('w:' + name), value)
borders.append(line)
p._p.get_or_add_pPr().append(borders)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
section.footer.paragraphs[0]._p.append(field)
doc.add_comment(run, text='Review this heading', author='Reviewer')
doc.add_paragraph('Body text preserved by the focused edit.')
doc.save('styled.docx')
edited = Document('styled.docx')
edited.paragraphs[1].runs[0].text = 'Body text updated by the focused edit.'
edited.save('styled.docx')
restored = Document('styled.docx')
# Word stores page margins in integer twips.
assert abs(restored.sections[0].top_margin - Mm(30)) < Mm(0.02)
assert restored.paragraphs[0].runs[0].font.size == Pt(22)
assert restored.paragraphs[0].paragraph_format.line_spacing == Pt(28)
assert next(iter(restored.comments)).text == 'Review this heading'
assert restored.sections[0].header.paragraphs[0].text == 'Office acceptance'
with ZipFile('styled.docx') as package:
    assert b'w:color="C00000"' in package.read('word/document.xml')
    assert b'w:eastAsia="Noto Sans CJK SC"' in package.read('word/document.xml')
    assert b'w:instr="PAGE"' in package.read('word/footer1.xml')

book = Workbook()
sheet = book.active
sheet.title = 'Budget'
for row in [('Item','Amount'),('Design',120),('Build',240),('Total','=SUM(B2:B3)')]:
    sheet.append(row)
for cell in sheet[1]:
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor='245A81')
sheet.freeze_panes = 'A2'
sheet.column_dimensions['A'].width = 24
chart = BarChart()
chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
sheet.add_chart(chart, 'D2')
validation = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1=0)
sheet.add_data_validation(validation)
validation.add('B2:B3')
book.create_sheet('Unchanged')['A1'] = 'Keep me'
book.save('styled.xlsx')
book = load_workbook('styled.xlsx')
book['Budget']['B2'] = 180
book.save('styled.xlsx')
book = load_workbook('styled.xlsx')
assert book['Budget']['B4'].value == '=SUM(B2:B3)'
assert book['Budget']['A1'].font.bold
assert book['Budget'].freeze_panes == 'A2'
assert len(book['Budget']._charts) == 1
assert book['Budget'].data_validations.count == 1
assert book['Unchanged']['A1'].value == 'Keep me'
with xlsxwriter.Workbook('cached.xlsx') as workbook:
    sheet = workbook.add_worksheet('Totals')
    sheet.write_number('A1', 3)
    sheet.write_number('A2', 4)
    sheet.write_formula('A3', '=SUM(A1:A2)', None, 7)
assert load_workbook('cached.xlsx', data_only=True)['Totals']['A3'].value == 7

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
text = slide.shapes.add_textbox(Inches(.6), Inches(.4), Inches(12), Inches(.8))
run = text.text_frame.paragraphs[0].add_run()
run.text, run.font.size = 'Quarterly results', SlidePt(32)
data = CategoryChartData()
data.categories = ['Q1', 'Q2', 'Q3']
data.add_series('Revenue', [100, 140, 180])
slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(10), Inches(4.5), data)
slide.notes_slide.notes_text_frame.text = 'Explain the growth trend.'
prs.save('styled.pptx')
prs = Presentation('styled.pptx')
prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].text = 'Updated results'
prs.save('styled.pptx')
prs = Presentation('styled.pptx')
assert prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size == SlidePt(32)
assert list(prs.slides[0].shapes[1].chart.series[0].values) == [100, 140, 180]
assert 'growth trend' in prs.slides[0].notes_slide.notes_text_frame.text

c = canvas.Canvas('form.pdf', pagesize=A4)
c.setFont('Helvetica', 18)
c.drawString(54, 780, 'Acceptance report')
c.acroForm.textfield(name='reviewer', x=54, y=700, width=200, height=24)
c.showPage()
c.save()
writer = PdfWriter(clone_from='form.pdf')
writer.update_page_form_field_values(writer.pages[0], {'reviewer':'Accepted'}, auto_regenerate=False)
writer.write('filled.pdf')
assert PdfReader('filled.pdf').get_fields()['reviewer']['/V'] == 'Accepted'
with pdfplumber.open('filled.pdf') as pdf:
    assert 'Acceptance report' in pdf.pages[0].extract_text()
with pdfium.PdfDocument('filled.pdf') as pdf:
    page = pdf[0]
    bitmap = page.render(scale=1)
    assert bitmap.width > 500 and bitmap.height > 800
    bitmap.close()
    page.close()
print('Verified DOCX styles/fields/comments, XLSX formulas/charts/validation, PPTX editable chart/notes, PDF forms/rendering')
